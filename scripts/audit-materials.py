"""Read-only source extraction for the prototype coverage review."""
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

root = Path(__file__).resolve().parents[1]
source = root / 'materials' / 'CRM-Counting Stars2.xlsx'
book = load_workbook(source, read_only=True, data_only=False)
result = {}
lines = []
for sheet in book:
    cells = []
    lines.append(f'\n## {sheet.title}')
    for row_number, row in enumerate(sheet, 1):
        values = []
        for cell in row:
            if cell.value is not None:
                text = str(cell.value).strip()
                if text:
                    cells.append({'cell': cell.coordinate, 'text': text})
                    values.append(f'{get_column_letter(cell.column)}: {text.replace(chr(10), " / ")}')
        if values:
            lines.append(f'{row_number} | ' + ' | '.join(values))
    result[sheet.title] = cells
assert len(result) == 19 and result['确认内容']
out = root / 'docs' / 'materials'
(out / 'workbook-audit.json').write_text(json.dumps({'source':source.name,'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'sheets':result},ensure_ascii=False,indent=2),encoding='utf-8')
(out / 'workbook-audit.txt').write_text('\n'.join(lines),encoding='utf-8')
print(json.dumps({name:len(cells) for name,cells in result.items()},ensure_ascii=False))
